import csv
import os

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

COLUMNS = ["Sourcerecord ID", "Source Title", "ISSN", "EISSN", "Active or Inactive", "Source Type"]


class Command(BaseCommand):
    help = "Extracts the main sheet of the Scopus source title list (xlsx) into Publications/data/scopus_sources.csv."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Path to the Scopus source title list downloaded from Elsevier.")
        parser.add_argument("--output", default=None, help="Where to write the csv file.")

    def handle(self, *args, **options):
        source = options["xlsx"]
        if not os.path.exists(source):
            raise CommandError("File not found: %s" % source)
        output = options["output"] or os.path.join(settings.BASE_DIR, "Publications", "data", "scopus_sources.csv")
        os.makedirs(os.path.dirname(output), exist_ok=True)

        workbook = load_workbook(source, read_only=True, data_only=True)
        written = 0
        for sheet in workbook.worksheets:
            indexes = None
            rows = []
            for row in sheet.iter_rows(values_only=True, max_col=16):
                names = [str(cell).strip() if cell else "" for cell in row]
                if indexes is None:
                    if all(name in names for name in COLUMNS):
                        indexes = [names.index(name) for name in COLUMNS]
                        continue
                    break
                if row[indexes[1]]:
                    rows.append(["" if row[index] is None else str(row[index]).strip() for index in indexes])
            if indexes is not None:
                with open(output, "w", encoding="utf-8", newline="") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(COLUMNS)
                    writer.writerows(rows)
                written = len(rows)
                break
        workbook.close()
        if not written:
            raise CommandError("No sheet with the columns %s was found." % ", ".join(COLUMNS))

        from Publications import verify

        verify._load_scopus_issns.cache_clear()
        self.stdout.write("Wrote %d sources to %s" % (written, output))
